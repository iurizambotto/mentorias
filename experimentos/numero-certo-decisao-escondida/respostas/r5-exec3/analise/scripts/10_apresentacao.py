"""Passo 10: monta entregas/apresentacao.pptx, 16:9, a partir do livro de numeros.

Uma cor de destaque so. O cinza carrega o contexto, o azul carrega o dado que o
titulo do slide aponta.
"""

from __future__ import annotations

import sys
from pathlib import Path

import matplotlib
import pandas as pd

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Inches, Pt

sys.path.insert(0, str(Path(__file__).parent))
sys.path.insert(0, ".claude/skills/analise-de-dados/scripts")

from numeros import carregar as carregar_livro

SAIDA = Path("entregas/apresentacao.pptx")
GRAFICOS = Path("analise/graficos")
DETALHE = Path("analise/blocos_detalhe.csv")

DESTAQUE = "#2a78d6"
CINZA = "#b9b8b2"
CINZA_ESCURO = "#52514e"
TINTA = RGBColor(0x0B, 0x0B, 0x0B)
TINTA2 = RGBColor(0x52, 0x51, 0x4E)
AZUL = RGBColor(0x2A, 0x78, 0xD6)
MUDO = RGBColor(0x89, 0x87, 0x81)

L, A = Inches(13.333), Inches(7.5)
MARGEM = Inches(0.72)
LARGURA_UTIL = L - 2 * MARGEM


def pct(v: float, casas: int = 2) -> str:
    return f"{v * 100:,.{casas}f}".replace(",", " ").replace(".", ",") + "%"


def num(v: float, casas: int = 0) -> str:
    return f"{v:,.{casas}f}".replace(",", "X").replace(".", ",").replace("X", ".")


def reais(v: float) -> str:
    return "R$ " + num(v, 2)


# ---------------------------------------------------------------- graficos
def estilo(ax) -> None:
    for lado in ["top", "right"]:
        ax.spines[lado].set_visible(False)
    ax.spines["left"].set_color("#c3c2b7")
    ax.spines["bottom"].set_color("#c3c2b7")
    ax.tick_params(colors=CINZA_ESCURO, labelsize=13)


def salvar(fig, nome: str) -> str:
    GRAFICOS.mkdir(parents=True, exist_ok=True)
    caminho = GRAFICOS / nome
    fig.savefig(caminho, dpi=200, bbox_inches="tight", transparent=True)
    plt.close(fig)
    return str(caminho)


def g_pacing_geral(v: dict) -> str:
    fig, ax = plt.subplots(figsize=(8.4, 3.5))
    rotulos = ["Investimento", "Impressões", "Cliques"]
    valores = [v["pacing.investimento"], v["pacing.impressoes"], v["pacing.cliques"]]
    cores = [CINZA, CINZA, DESTAQUE]  # o titulo do slide aponta o clique
    barras = ax.barh(rotulos, valores, color=cores, height=0.56)
    ax.axvline(1.0, color=CINZA_ESCURO, linestyle="--", linewidth=1.4)
    # acima do plot, em coordenadas do eixo, para nao cair sobre o rotulo de 100%
    ax.annotate("meta do plano", xy=(1.0, 1.0), xycoords=("data", "axes fraction"),
                xytext=(0, 8), textcoords="offset points",
                ha="center", fontsize=12, color=CINZA_ESCURO)
    for barra, valor in zip(barras, valores):
        ax.text(valor + 0.025, barra.get_y() + barra.get_height() / 2, pct(valor, 1),
                va="center", fontsize=14, fontweight="bold", color="#0b0b0b")
    ax.set_xlim(0, 1.45)
    ax.invert_yaxis()
    ax.set_xticks([0, 0.5, 1.0])
    ax.set_xticklabels(["0%", "50%", "100%"])
    estilo(ax)
    return salvar(fig, "pacing_geral.png")


def g_par(planejado: float, realizado: float, formatar, nome: str) -> str:
    fig, ax = plt.subplots(figsize=(5.0, 3.5))
    barras = ax.bar(["planejado", "realizado"], [planejado, realizado],
                    color=[CINZA, DESTAQUE], width=0.52)
    for barra, valor in zip(barras, [planejado, realizado]):
        ax.text(barra.get_x() + barra.get_width() / 2, valor * 1.03, formatar(valor),
                ha="center", fontsize=16, fontweight="bold", color="#0b0b0b")
    ax.set_ylim(0, max(planejado, realizado) * 1.25)
    ax.set_yticks([])
    ax.spines["left"].set_visible(False)
    estilo(ax)
    return salvar(fig, nome)


def g_veiculos(v: dict) -> str:
    fig, ax = plt.subplots(figsize=(9.2, 3.7))
    veiculos = ["Meta Ads", "Youtube Ads", "Tiktok Ads"]
    chaves = ["meta", "youtube", "tiktok"]
    metricas = ["investimento", "impressoes", "cliques"]
    nomes = ["Investimento", "Impressões", "Cliques"]
    largura = 0.26
    posicoes = range(len(veiculos))
    for i, (metrica, nome) in enumerate(zip(metricas, nomes)):
        valores = [v[f"veiculo.{k}.pacing_{metrica}"] for k in chaves]
        # so o TikTok recebe o destaque: e dele que o titulo fala
        cores = [DESTAQUE if k == "tiktok" else CINZA for k in chaves]
        deslocamento = [p + (i - 1) * largura for p in posicoes]
        barras = ax.bar(deslocamento, valores, largura * 0.9, color=cores)
        for barra, valor in zip(barras, valores):
            ax.text(barra.get_x() + barra.get_width() / 2, valor + 0.035, pct(valor, 0),
                    ha="center", fontsize=11.5, color="#0b0b0b", fontweight="bold")
            # o nome da metrica vai girado sob cada barra: na horizontal os tres
            # rotulos de um mesmo grupo se sobrepoem
            ax.annotate(nome, xy=(barra.get_x() + barra.get_width() / 2, 0),
                        xycoords=("data", "axes fraction"), xytext=(0, -6),
                        textcoords="offset points", rotation=90,
                        ha="center", va="top", fontsize=10, color=CINZA_ESCURO)
    ax.axhline(1.0, color=CINZA_ESCURO, linestyle="--", linewidth=1.4)
    ax.set_xticks(list(posicoes))
    ax.set_xticklabels(veiculos, fontsize=14)
    ax.tick_params(axis="x", pad=74, length=0)  # espaco para os rotulos girados
    ax.set_yticks([0, 0.5, 1.0, 1.5])
    ax.set_yticklabels(["0%", "50%", "100%", "150%"])
    ax.set_ylim(0, 1.62)
    estilo(ax)
    return salvar(fig, "veiculos.png")


def g_campanhas() -> str:
    detalhe = pd.read_csv(DETALHE).dropna(subset=["pacing_impressoes"])
    agrupado = detalhe.groupby("Campanha")[["plan_Soma de Impressoes", "real_Soma de Impressoes"]].sum()
    agrupado["pacing"] = agrupado["real_Soma de Impressoes"] / agrupado["plan_Soma de Impressoes"]
    serie = agrupado["pacing"].sort_values()

    fig, ax = plt.subplots(figsize=(9.0, 5.1))
    # destaque nos extremos: a que nao foi ao ar e a que quase triplicou a meta
    cores = [DESTAQUE if i in (0, len(serie) - 1) else CINZA for i in range(len(serie))]
    barras = ax.barh(serie.index, serie.values, color=cores, height=0.66)
    ax.axvline(1.0, color=CINZA_ESCURO, linestyle="--", linewidth=1.4)
    for barra, valor in zip(barras, serie.values):
        ax.text(valor + 0.04, barra.get_y() + barra.get_height() / 2, pct(valor, 0),
                va="center", fontsize=11, color="#0b0b0b")
    ax.set_xlim(0, 3.25)
    ax.set_xticks([0, 1.0, 2.0, 3.0])
    ax.set_xticklabels(["0%", "100%", "200%", "300%"])
    ax.tick_params(axis="y", labelsize=11.5)
    estilo(ax)
    return salvar(fig, "campanhas.png")


def g_cobertura(v: dict) -> str:
    fig, ax = plt.subplots(figsize=(9.0, 1.9))
    dentro = v["real.total_investimento"]
    fora = v["real.investimento_total_arquivo"] - dentro
    ax.barh([""], [dentro], color=DESTAQUE, height=0.52)
    ax.barh([""], [fora], left=[dentro], color=CINZA, height=0.52)
    # a fatia com plano e estreita demais para receber o rotulo dentro
    ax.annotate(f"com plano\n{pct(v['real.share_investimento_no_pacing'], 2)}",
                xy=(dentro, 0.26), xytext=(dentro * 1.6, 0.75),
                ha="center", fontsize=17, fontweight="bold", color=DESTAQUE,
                arrowprops=dict(arrowstyle="-", color=DESTAQUE, linewidth=1.4))
    ax.text(dentro + fora / 2, 0, "sem nenhuma linha de plano", ha="center", va="center",
            fontsize=18, fontweight="bold", color="#0b0b0b")
    ax.set_ylim(-0.4, 1.25)
    ax.set_xlim(0, v["real.investimento_total_arquivo"])
    ax.set_xticks([])
    ax.set_yticks([])
    for lado in ["top", "right", "left", "bottom"]:
        ax.spines[lado].set_visible(False)
    return salvar(fig, "cobertura.png")


# ---------------------------------------------------------------- slides
def caixa(slide, texto: str, esquerda, topo, largura, altura, tamanho: int,
          *, cor=TINTA, negrito=False, alinhamento=PP_ALIGN.LEFT, espacamento: float = 1.12):
    forma = slide.shapes.add_textbox(esquerda, topo, largura, altura)
    quadro = forma.text_frame
    quadro.word_wrap = True
    for i, linha in enumerate(texto.split("\n")):
        paragrafo = quadro.paragraphs[0] if i == 0 else quadro.add_paragraph()
        paragrafo.text = linha
        paragrafo.alignment = alinhamento
        paragrafo.line_spacing = espacamento
        for corrida in paragrafo.runs:
            corrida.font.size = Pt(tamanho)
            corrida.font.bold = negrito
            corrida.font.color.rgb = cor
            corrida.font.name = "Calibri"
    return forma


def limpar_placeholders(slide, manter: set[int]) -> None:
    """Remove placeholder que o layout traz e o slide nao usa: data, rodape, numero."""
    for forma in list(slide.placeholders):
        if forma.placeholder_format.idx not in manter:
            forma._element.getparent().remove(forma._element)


def novo_slide(prs: Presentation):
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # layout em branco
    limpar_placeholders(slide, manter=set())
    return slide


def slide_titulo(prs, titulo: str, subtitulo: str = "") -> object:
    """Usa o layout 'Title Only' para que o titulo seja um placeholder de verdade,
    e nao uma caixa de texto que leitores de tela e o validador nao reconhecem."""
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    limpar_placeholders(slide, manter={0})
    forma = slide.shapes.title
    forma.left, forma.top = MARGEM, Inches(0.46)
    forma.width, forma.height = LARGURA_UTIL, Inches(1.0)
    quadro = forma.text_frame
    quadro.word_wrap = True
    quadro.text = titulo
    paragrafo = quadro.paragraphs[0]
    paragrafo.alignment = PP_ALIGN.LEFT
    paragrafo.line_spacing = 1.08
    # o titulo tem uma linha so: acima de ~48 caracteres ele quebraria e invadiria
    # a regua e o subtitulo, entao o corpo diminui em vez de a caixa crescer
    tamanho = 30 if len(titulo) <= 48 else 26
    for corrida in paragrafo.runs:
        corrida.font.size = Pt(tamanho)
        corrida.font.bold = True
        corrida.font.color.rgb = TINTA
        corrida.font.name = "Calibri"
    if subtitulo:
        caixa(slide, subtitulo, MARGEM, Inches(1.36), LARGURA_UTIL, Inches(0.6), 17, cor=TINTA2)
    return slide


def regua(slide, topo=Inches(1.28)) -> None:
    from pptx.enum.shapes import MSO_SHAPE

    forma = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, MARGEM, topo, Inches(1.5), Emu(26000))
    forma.fill.solid()
    forma.fill.fore_color.rgb = AZUL
    forma.line.fill.background()
    forma.shadow.inherit = False


def imagem_centrada(slide, caminho: str, topo, altura) -> None:
    from PIL import Image

    with Image.open(caminho) as img:
        proporcao = img.width / img.height
    largura = Emu(int(altura * proporcao))
    if largura > LARGURA_UTIL:
        largura = LARGURA_UTIL
        altura = Emu(int(largura / proporcao))
    slide.shapes.add_picture(caminho, Emu(int((L - largura) / 2)), topo, largura, altura)


def main() -> int:
    livro = carregar_livro()
    v = {chave: item["valor"] for chave, item in livro.items()}

    prs = Presentation()
    prs.slide_width, prs.slide_height = L, A

    # 1. capa
    slide = novo_slide(prs)
    caixa(slide, "O dinheiro foi entregue.\nO clique não.", MARGEM, Inches(2.1),
          Inches(10.4), Inches(2.4), 46, negrito=True, espacamento=1.06)
    caixa(slide, "Pacing das campanhas de mídia  ·  plano de 28/05/2024 a 31/07/2024",
          MARGEM, Inches(4.5), Inches(10.4), Inches(0.5), 19, cor=TINTA2)
    caixa(slide, f"{num(v['plan.flights'])} flights planejados  ·  "
                 f"{num(v['real.linhas'])} linhas de entrega diária  ·  "
                 f"{num(v['real.linhas_no_pacing'])} entram na conta",
          MARGEM, Inches(5.12), Inches(11.0), Inches(0.5), 15, cor=MUDO)

    # 2. a pergunta
    slide = slide_titulo(prs, "A pergunta")
    regua(slide)
    caixa(slide, "“A gente entregou o que tinha planejado?”", MARGEM, Inches(2.3),
          Inches(11.2), Inches(1.2), 34, negrito=True)
    caixa(slide,
          "Pacing é o realizado dividido pelo planejado.\n"
          "Só entra na conta a entrega que caiu dentro da janela do flight, no mesmo veículo.\n"
          "Campanha sem linha de plano fica de fora.",
          MARGEM, Inches(3.75), Inches(11.2), Inches(1.8), 18, cor=TINTA2, espacamento=1.5)

    # 3. a resposta
    slide = slide_titulo(prs, "A resposta: depende de qual das três métricas você olha")
    regua(slide)
    caixa(slide, "A verba foi cumprida. A audiência veio acima. O clique veio pela metade.",
          MARGEM, Inches(1.62), LARGURA_UTIL, Inches(0.6), 20, cor=TINTA2)
    kpis = [
        ("Investimento", pct(v["pacing.investimento"]), "no plano"),
        ("Impressões", pct(v["pacing.impressoes"]), "acima do plano"),
        ("Cliques", pct(v["pacing.cliques"]), "abaixo do plano"),
    ]
    largura_kpi = Inches(3.72)
    for i, (rotulo, valor, nota) in enumerate(kpis):
        esquerda = MARGEM + i * (largura_kpi + Inches(0.26))
        caixa(slide, rotulo.upper(), esquerda, Inches(2.75), largura_kpi, Inches(0.42), 15, cor=MUDO)
        caixa(slide, valor, esquerda, Inches(3.2), largura_kpi, Inches(1.1), 50, negrito=True,
              cor=AZUL if i == 2 else TINTA)
        caixa(slide, nota, esquerda, Inches(4.42), largura_kpi, Inches(0.5), 17, cor=TINTA2)
    caixa(slide, f"Base: {livro['pacing.impressoes']['linhas']} blocos de plano e "
                 f"{num(v['real.linhas_no_pacing'])} linhas de entrega dentro das janelas.",
          MARGEM, Inches(5.5), LARGURA_UTIL, Inches(0.5), 14, cor=MUDO)

    # 4. insight: a verba nao responde a pergunta
    slide = slide_titulo(prs, "Cumprir a verba não é entregar o plano",
                         f"Faltaram {num(v['gap.cliques'])} cliques que estavam no plano.")
    regua(slide)
    imagem_centrada(slide, g_pacing_geral(v), Inches(2.0), Inches(3.9))
    caixa(slide, "Quem olhar só a linha de orçamento vai concluir que a campanha correu bem.",
          MARGEM, Inches(6.2), LARGURA_UTIL, Inches(0.5), 17, cor=TINTA2)

    # 5. insight: CPM
    slide = slide_titulo(prs, f"A mídia veio {pct(v['var.queda_cpm'], 2)} mais barata que o plano supunha",
                         "Custo por mil impressões, planejado contra realizado.")
    regua(slide)
    imagem_centrada(slide, g_par(v["plan.cpm"], v["real.cpm"], reais, "cpm.png"), Inches(1.95), Inches(3.9))
    caixa(slide, f"Com o mesmo dinheiro, isso gerou {num(v['gap.impressoes'])} impressões além do plano. "
                 f"É a parte boa do resultado, e também a explicação do próximo slide.",
          MARGEM, Inches(6.1), LARGURA_UTIL, Inches(0.8), 17, cor=TINTA2)

    # 6. insight: CTR
    slide = slide_titulo(prs, f"E a taxa de clique veio {pct(v['var.queda_ctr'], 2)} abaixo",
                         "Percentual de quem viu o anúncio e clicou, planejado contra realizado.")
    regua(slide)
    imagem_centrada(slide, g_par(v["plan.ctr"], v["real.ctr"], lambda x: pct(x, 2), "ctr.png"),
                    Inches(1.95), Inches(3.9))
    caixa(slide, "A mídia barata que trouxe a audiência extra é a mesma que converte menos em clique. "
                 "Não são dois problemas, é um só, visto de dois lados.",
          MARGEM, Inches(6.1), LARGURA_UTIL, Inches(0.8), 17, cor=TINTA2)

    # 7. insight: veiculos
    slide = slide_titulo(prs, "O TikTok ficou abaixo da meta nas três métricas",
                         "Meta trouxe volume sem clique. YouTube trouxe clique sem volume.")
    regua(slide)
    imagem_centrada(slide, g_veiculos(v), Inches(1.95), Inches(4.0))
    caixa(slide, f"Sobraram {reais(v['veiculo.tiktok.verba_nao_gasta'])} de verba planejada e não gasta no veículo. "
                 f"A ação Inauguração tinha {reais(v['inauguracao.investimento_planejado'])} reservados e não teve um dia de entrega.",
          MARGEM, Inches(6.15), LARGURA_UTIL, Inches(0.8), 16, cor=TINTA2)

    # 8. insight: campanhas
    slide = slide_titulo(prs, "De zero ao triplo da meta, dependendo da campanha",
                         "Pacing de impressões por campanha.")
    regua(slide)
    imagem_centrada(slide, g_campanhas(), Inches(1.8), Inches(4.7))
    caixa(slide, "O número consolidado esconde essa dispersão. A conversa útil é campanha a campanha.",
          MARGEM, Inches(6.62), LARGURA_UTIL, Inches(0.5), 16, cor=TINTA2)

    # 9. cobertura
    slide = slide_titulo(prs, f"Este plano cobre {pct(v['real.share_investimento_no_pacing'], 2)} do que a área gastou",
                         f"{num(v['real.campanhas_sem_plano'])} campanhas rodaram sem nenhuma linha de plano.")
    regua(slide)
    imagem_centrada(slide, g_cobertura(v), Inches(2.4), Inches(2.1))
    caixa(slide, f"Investimento no arquivo: {reais(v['real.investimento_total_arquivo'])}.  "
                 f"Coberto por este plano: {reais(v['real.total_investimento'])}.",
          MARGEM, Inches(4.85), LARGURA_UTIL, Inches(0.5), 17, cor=TINTA2)
    caixa(slide, "Os números deste deck são verdadeiros sobre o plano. Eles não descrevem a operação inteira.",
          MARGEM, Inches(5.5), LARGURA_UTIL, Inches(0.6), 17, cor=TINTA)

    # 10. o que nao responde e qualidade da base
    slide = slide_titulo(prs, "O que esta análise não responde")
    regua(slide)
    caixa(slide,
          "Resultado de negócio: o arquivo tem mídia, não venda.\n"
          "Qual métrica era a meta de cada flight.\n"
          "O motivo da queda de engajamento.\n"
          "Pacing de flight isolado em três janelas sobrepostas.",
          MARGEM, Inches(1.8), LARGURA_UTIL, Inches(2.3), 18, cor=TINTA2, espacamento=1.8)
    caixa(slide, "Qualidade da base", MARGEM, Inches(4.3), LARGURA_UTIL, Inches(0.5), 21, negrito=True)
    caixa(slide,
          "Extensão .xls, conteúdo CSV, e duas tabelas empilhadas.\n"
          "Duas reservas sem meta no plano: ficaram fora da conta.\n"
          f"Público vem como “N/a” em {num(v['qualidade.publico_nulo_em_texto'])} linhas.",
          MARGEM, Inches(4.9), LARGURA_UTIL, Inches(1.5), 17, cor=TINTA2, espacamento=1.8)

    # 11. decisoes de metodo
    slide = slide_titulo(prs, "Duas decisões de método mudam o resultado",
                         "Cada alternativa foi calculada, não estimada.")
    regua(slide)
    tabela = slide.shapes.add_table(3, 3, MARGEM, Inches(1.95), LARGURA_UTIL, Inches(3.0)).table
    tabela.columns[0].width = Emu(int(LARGURA_UTIL * 0.42))
    tabela.columns[1].width = Emu(int(LARGURA_UTIL * 0.29))
    tabela.columns[2].width = Emu(int(LARGURA_UTIL * 0.29))
    conteudo = [
        ["Decisão", "Como foi feito", "Se fosse do outro jeito"],
        ["O mesmo dia de entrega não foi contado duas vezes",
         f"Impressões: {pct(v['pacing.impressoes'])}",
         f"{pct(v['alt.pacing_sem_dedup_impressoes'])}, ou {pct(v['alt.inflacao_sem_dedup_impressoes'])} a mais de entrega que nunca existiu"],
        ["Somamos tudo antes de dividir, em vez de tirar média entre reservas",
         f"Cliques: {pct(v['pacing.cliques'])}",
         f"{pct(v['alt.media_simples_cliques'])}, e a campanha pareceria ter superado a meta"],
    ]
    for i, linha in enumerate(conteudo):
        for j, texto in enumerate(linha):
            celula = tabela.cell(i, j)
            celula.text = texto
            celula.margin_left = celula.margin_right = Inches(0.13)
            celula.margin_top = celula.margin_bottom = Inches(0.09)
            for paragrafo in celula.text_frame.paragraphs:
                paragrafo.line_spacing = 1.15
                for corrida in paragrafo.runs:
                    corrida.font.size = Pt(14)
                    corrida.font.bold = i == 0
                    corrida.font.color.rgb = TINTA if i == 0 else TINTA2
                    corrida.font.name = "Calibri"
    caixa(slide, "As demais decisões, com o valor de cada alternativa, estão em analise/decisoes.md.",
          MARGEM, Inches(5.25), LARGURA_UTIL, Inches(0.5), 15, cor=MUDO)

    # 12. proximos passos
    slide = slide_titulo(prs, "Próximos passos")
    regua(slide)
    caixa(slide,
          "1.  Definir qual métrica é a meta de cada flight.\n"
          f"2.  Responder o caso do TikTok: {reais(v['inauguracao.investimento_planejado'])} reservados e nenhum dia no ar.\n"
          f"3.  Reconciliar as datas do plano: {reais(v['real.investimento_fora_da_janela'])} entregues fora da janela.\n"
          f"4.  Decidir o que fazer com as {num(v['real.campanhas_sem_plano'])} campanhas sem plano.\n"
          f"5.  Revisar a premissa de preço: o plano usou {reais(v['plan.cpm'])} por mil e o mercado entregou a {reais(v['real.cpm'])}.",
          MARGEM, Inches(1.85), LARGURA_UTIL, Inches(3.6), 19, cor=TINTA2, espacamento=1.95)

    # 13. apendice
    slide = slide_titulo(prs, "Apêndice: como reproduzir")
    regua(slide)
    caixa(slide,
          "perfilar.py sobre o arquivo, depois os scripts numerados de analise/scripts/, na ordem.\n\n"
          "Cada número deste deck está em analise/numeros.json com o script que o calculou.\n"
          "As decisões e o resultado de cada alternativa estão em analise/decisoes.md.\n"
          "O detalhe por bloco de plano está em analise/blocos_detalhe.csv e no dashboard.",
          MARGEM, Inches(1.9), LARGURA_UTIL, Inches(3.0), 18, cor=TINTA2, espacamento=1.6)

    SAIDA.parent.mkdir(parents=True, exist_ok=True)
    prs.save(SAIDA)
    print(f"gravado {SAIDA} ({len(prs.slides.__iter__.__self__._sldIdLst)} slides)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
