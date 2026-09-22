from pathlib import Path

from docx import Document
from md_para_docx import converter


def test_headings_bullets_table_and_bold(tmp_path: Path) -> None:
    md = tmp_path / "i.md"
    md.write_text(
        "# Titulo\n\nParagrafo com **negrito**.\n\n## Secao\n\n- item um\n- item dois\n\n"
        "1. passo\n\n| a | b |\n|---|---|\n| 1 | 2 |\n",
        encoding="utf-8",
    )
    saida = converter(md, tmp_path / "i.docx")
    doc = Document(str(saida))
    estilos = [(p.style.name, p.text) for p in doc.paragraphs if p.text]
    assert ("Title", "Titulo") in estilos or ("Heading 1", "Titulo") in estilos
    assert ("Heading 2", "Secao") in estilos
    assert ("List Bullet", "item um") in estilos
    assert ("List Number", "passo") in estilos
    assert any(r.bold and r.text == "negrito" for p in doc.paragraphs for r in p.runs)
    assert doc.tables[0].cell(1, 1).text == "2"
    assert doc.core_properties.title == "Titulo"


def test_pdf_usa_perfil_isolado_do_libreoffice(monkeypatch, tmp_path: Path) -> None:
    # Concurrent soffice calls clash over the shared user profile, so each run needs its own.
    import md_para_docx

    capturado: dict = {}

    def falso_run(comando, **kwargs):
        capturado["comando"] = comando
        (tmp_path / "x.pdf").write_bytes(b"%PDF-1.4")
        return None

    monkeypatch.setattr(md_para_docx.shutil, "which", lambda _: "/usr/bin/soffice")
    monkeypatch.setattr(md_para_docx.subprocess, "run", falso_run)
    docx = tmp_path / "x.docx"
    docx.write_bytes(b"fake")
    md_para_docx.para_pdf(docx)
    assert any(a.startswith("-env:UserInstallation=file://") for a in capturado["comando"])
