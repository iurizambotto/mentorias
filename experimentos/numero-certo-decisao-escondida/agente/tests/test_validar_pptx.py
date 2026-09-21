from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt
from validar_pptx import validar


def _deck(tmp_path: Path, *, fonte: int = 18, fora: bool = False, vazio: bool = False) -> Path:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    for i in range(6):
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        if not (vazio and i == 0):
            slide.shapes.title.text = f"Mensagem {i}"
        left = Inches(12.5) if (fora and i == 1) else Inches(1)
        caixa = slide.shapes.add_textbox(left, Inches(2), Inches(4), Inches(1))
        run = caixa.text_frame.paragraphs[0].add_run()
        run.text = "texto de apoio"
        run.font.size = Pt(fonte)
    saida = tmp_path / "d.pptx"
    prs.save(str(saida))
    return saida


def test_clean_deck_passes(tmp_path: Path) -> None:
    assert validar(_deck(tmp_path)) == []


def test_small_font_is_flagged(tmp_path: Path) -> None:
    assert any("fonte" in p for p in validar(_deck(tmp_path, fonte=9)))


def test_shape_out_of_bounds_is_flagged(tmp_path: Path) -> None:
    assert any("fora do slide" in p for p in validar(_deck(tmp_path, fora=True)))


def test_empty_placeholder_is_flagged(tmp_path: Path) -> None:
    assert any("vazio" in p for p in validar(_deck(tmp_path, vazio=True)))


def test_render_usa_perfil_isolado_do_libreoffice(monkeypatch, tmp_path: Path) -> None:
    import validar_pptx

    comandos: list[list[str]] = []

    def falso_run(comando, **kwargs):
        comandos.append(comando)
        return None

    monkeypatch.setattr(validar_pptx.shutil, "which", lambda nome: f"/usr/bin/{nome}")
    monkeypatch.setattr(validar_pptx.subprocess, "run", falso_run)
    validar_pptx.renderizar(_deck(tmp_path), tmp_path / "preview")
    assert any(a.startswith("-env:UserInstallation=file://") for a in comandos[0])
