"""Convert a small, predictable markdown subset into a clean DOCX.

Supported: # to ### headings, paragraphs, - bullets, 1. numbered items,
pipe tables, > quotes and **bold** inline. Optionally exports PDF through
LibreOffice when it is installed.
"""

from __future__ import annotations

import argparse
import re
import shutil
import subprocess
import tempfile
from pathlib import Path

from docx import Document
from docx.shared import Pt

_NEGRITO = re.compile(r"(\*\*[^*]+\*\*)")


def _texto_com_negrito(paragrafo, texto: str) -> None:
    for parte in _NEGRITO.split(texto):
        if not parte:
            continue
        if parte.startswith("**") and parte.endswith("**"):
            paragrafo.add_run(parte[2:-2]).bold = True
        else:
            paragrafo.add_run(parte)


def _tabela(doc: Document, linhas: list[str]) -> None:
    celulas = [[c.strip() for c in linha.strip().strip("|").split("|")] for linha in linhas]
    celulas = [c for c in celulas if not all(re.fullmatch(r":?-{2,}:?", x) for x in c)]
    tabela = doc.add_table(rows=len(celulas), cols=max(len(c) for c in celulas))
    tabela.style = "Table Grid"
    for i, linha in enumerate(celulas):
        for j, valor in enumerate(linha):
            celula = tabela.cell(i, j)
            celula.text = ""
            _texto_com_negrito(celula.paragraphs[0], valor)
            if i == 0:
                for run in celula.paragraphs[0].runs:
                    run.bold = True


def converter(origem: Path, destino: Path) -> Path:
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)
    linhas = origem.read_text(encoding="utf-8").splitlines()
    titulo_definido = False
    i = 0
    while i < len(linhas):
        linha = linhas[i]
        if not linha.strip():
            i += 1
            continue
        if linha.lstrip().startswith("|"):
            bloco = []
            while i < len(linhas) and linhas[i].lstrip().startswith("|"):
                bloco.append(linhas[i])
                i += 1
            _tabela(doc, bloco)
            continue
        cabecalho = re.match(r"^(#{1,3})\s+(.*)$", linha)
        if cabecalho:
            nivel, texto = len(cabecalho.group(1)), cabecalho.group(2).strip()
            if nivel == 1 and not titulo_definido:
                doc.add_heading(texto, level=0)
                doc.core_properties.title = texto
                titulo_definido = True
            else:
                doc.add_heading(texto, level=nivel)
        elif re.match(r"^\s*[-*]\s+", linha):
            _texto_com_negrito(doc.add_paragraph(style="List Bullet"), re.sub(r"^\s*[-*]\s+", "", linha))
        elif re.match(r"^\s*\d+[.)]\s+", linha):
            _texto_com_negrito(doc.add_paragraph(style="List Number"), re.sub(r"^\s*\d+[.)]\s+", "", linha))
        elif linha.startswith(">"):
            p = doc.add_paragraph()
            _texto_com_negrito(p, linha.lstrip("> ").strip())
            for run in p.runs:
                run.italic = True
        else:
            _texto_com_negrito(doc.add_paragraph(), linha.strip())
        i += 1
    destino.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(destino))
    return destino


def para_pdf(docx: Path) -> Path | None:
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        return None
    # A private profile per call, so two conversions can run at the same time.
    perfil = tempfile.mkdtemp(prefix="soffice-")
    try:
        subprocess.run(
            [
                soffice,
                f"-env:UserInstallation=file://{perfil}",
                "--headless",
                "--convert-to",
                "pdf",
                "--outdir",
                str(docx.parent),
                str(docx),
            ],
            check=True,
            capture_output=True,
            timeout=300,
        )
    finally:
        shutil.rmtree(perfil, ignore_errors=True)
    pdf = docx.with_suffix(".pdf")
    return pdf if pdf.exists() else None


def main() -> int:
    parser = argparse.ArgumentParser(description="Markdown simples para DOCX, e PDF opcional.")
    parser.add_argument("origem", type=Path)
    parser.add_argument("destino", type=Path)
    parser.add_argument("--pdf", action="store_true")
    args = parser.parse_args()
    saida = converter(args.origem, args.destino)
    print(f"docx: {saida}")
    if args.pdf:
        pdf = para_pdf(saida)
        print(f"pdf: {pdf}" if pdf else "pdf: LibreOffice indisponivel, PDF nao gerado")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
