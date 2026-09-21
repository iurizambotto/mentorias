"""Mechanical checks on a results deck, plus an optional render to PNG.

Catches what a model cannot see in its own python-pptx code: text too small to
read, shapes that leave the slide, leftover empty placeholders and slides with
too many words. Rendering lets the author look at every slide before delivery.
"""

from __future__ import annotations

import argparse
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from pptx import Presentation
from pptx.util import Emu, Pt

_FONTE_MINIMA = Pt(12)
_PALAVRAS_MAXIMAS = 70
_FOLGA = Emu(45720)


def validar(caminho: Path, minimo: int = 5, maximo: int = 16) -> list[str]:
    prs = Presentation(str(caminho))
    largura, altura = prs.slide_width, prs.slide_height
    problemas: list[str] = []
    total = len(prs.slides)
    if not minimo <= total <= maximo:
        problemas.append(f"{total} slides, esperado entre {minimo} e {maximo}")

    for n, slide in enumerate(prs.slides, start=1):
        palavras = 0
        titulo = slide.shapes.title.text.strip() if slide.shapes.title is not None else ""
        for forma in slide.shapes:
            if (
                forma.left is not None
                and forma.width is not None
                and (
                    forma.left < -_FOLGA
                    or forma.top < -_FOLGA
                    or forma.left + forma.width > largura + _FOLGA
                    or forma.top + forma.height > altura + _FOLGA
                )
            ):
                problemas.append(f"slide {n}: '{forma.name}' sai fora do slide")
            if forma.is_placeholder and forma.has_text_frame and not forma.text_frame.text.strip():
                problemas.append(f"slide {n}: placeholder vazio '{forma.name}'")
            if not forma.has_text_frame:
                continue
            if forma != slide.shapes.title:
                palavras += len(forma.text_frame.text.split())
            for paragrafo in forma.text_frame.paragraphs:
                for run in paragrafo.runs:
                    if run.font.size is not None and run.font.size < _FONTE_MINIMA and run.text.strip():
                        problemas.append(f"slide {n}: fonte de {run.font.size.pt:.0f} pt em '{run.text[:30]}'")
        if not titulo and n > 1:
            problemas.append(f"slide {n}: sem titulo")
        if palavras > _PALAVRAS_MAXIMAS:
            problemas.append(f"slide {n}: {palavras} palavras de corpo, acima de {_PALAVRAS_MAXIMAS}")
    return problemas


def renderizar(caminho: Path, pasta: Path) -> list[Path]:
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice or not shutil.which("pdftoppm"):
        return []
    pasta.mkdir(parents=True, exist_ok=True)
    # A private profile per call, so two renders can run at the same time.
    perfil = tempfile.mkdtemp(prefix="soffice-")
    try:
        subprocess.run(
            [
                soffice,
                f"-env:UserInstallation=file://{perfil}",
                "--headless",
                "--convert-to",
                "pdf",
                "--outdir",
                str(pasta),
                str(caminho),
            ],
            check=True,
            capture_output=True,
            timeout=600,
        )
    finally:
        shutil.rmtree(perfil, ignore_errors=True)
    pdf = pasta / caminho.with_suffix(".pdf").name
    subprocess.run(
        ["pdftoppm", "-png", "-r", "60", str(pdf), str(pasta / "slide")], check=True, capture_output=True, timeout=300
    )
    return sorted(pasta.glob("slide*.png"))


def main() -> int:
    parser = argparse.ArgumentParser(description="Valida um deck PPTX e opcionalmente renderiza os slides.")
    parser.add_argument("arquivo", type=Path)
    parser.add_argument("--render", type=Path, help="pasta para os PNG de cada slide")
    args = parser.parse_args()
    problemas = validar(args.arquivo)
    for p in problemas:
        print(f"PROBLEMA: {p}")
    print("ok: deck sem problema mecanico" if not problemas else f"{len(problemas)} problema(s)")
    if args.render:
        imagens = renderizar(args.arquivo, args.render)
        print(
            "\n".join(f"render: {i}" for i in imagens)
            if imagens
            else "render indisponivel (LibreOffice ou pdftoppm ausente)"
        )
    return 1 if problemas else 0


if __name__ == "__main__":
    sys.exit(main())
